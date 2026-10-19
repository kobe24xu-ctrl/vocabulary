#!/usr/bin/env python3
"""Import the five master-owned rosters and class courses from a feedback workbook.

Preview: python scripts/import-feedback.py /private/feedback.xlsx
Apply:   python scripts/import-feedback.py /private/feedback.xlsx --apply --password-stdin
The initial password is read from stdin, hashed per student, and never logged.
Existing activated passwords and historical student IDs are preserved on rerun.
"""
import argparse
import collections
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import sys
import time
import unicodedata
import uuid

SUBJECTS = {'数学':'Mathematics','英语':'English','经济':'Economics','会计':'Accounting','计算机':'Computer Science','学术英语':'Academic English','语言学习':'Language Learning','商务':'Business','物理':'Physics','化学':'Chemistry'}
SHEETS = {'新南总表':('UNSW 新南预科',None),'沃隆港总表':('UOW 伍伦贡预科',None),'NCUK1班总表':('NCUK 国际预科',1),'NCUK2班总表':('NCUK 国际预科',2),'NCUK3班总表':('NCUK 国际预科',3)}
ROOT=Path(__file__).resolve().parent.parent


def parse_workbook(path):
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError('Install openpyxl==3.1.5 for workbook imports.') from None
    workbook=openpyxl.load_workbook(path,data_only=True)
    parsed=[]
    for title,(program,number) in SHEETS.items():
        if title not in workbook.sheetnames:
            raise RuntimeError('Expected feedback summary sheet is missing: '+title)
        sheet=workbook[title]
        if str(sheet.cell(3,2).value).strip()!='学生姓名':
            raise RuntimeError('Unexpected roster header in '+title)
        names=[]
        rows=[]
        for row in range(5,sheet.max_row+1):
            if isinstance(sheet.cell(row,1).value,(int,float)) and sheet.cell(row,2).value:
                name=str(sheet.cell(row,2).value).strip()
                username=unicodedata.normalize('NFKC',name).strip().lower()
                if not username or len(username)>60 or username=='master':
                    raise RuntimeError('A roster name cannot be used as a login; no import performed.')
                names.append((name,username)); rows.append(row)
        if not names:
            raise RuntimeError('Empty roster in '+title)
        courses=[]
        for column in range(3,sheet.max_column+1,7):
            raw=sheet.cell(3,column).value
            if raw is None:
                continue
            subject=str(raw).strip()
            if subject not in SUBJECTS:
                raise RuntimeError('Unmapped course: '+subject)
            contents=[]
            for row in rows:
                cell=sheet.cell(row,column+4).value
                if cell is not None:
                    text=str(cell).strip()
                    if text and text not in contents:
                        contents.append(text)
            content='\n\n'.join(contents)
            if len(content)>10000:
                raise RuntimeError('Course outline exceeds supported length in '+title)
            courses.append({'subject':SUBJECTS[subject],'title':subject,'content':content,'source':path.name+' · '+title})
        name=str(sheet.cell(1,1).value).split(' · ')[0].strip()
        if not name or len(name)>100:
            raise RuntimeError('Invalid class title in '+title)
        parsed.append({'sheet':title,'program':program,'number':number,'name':name,'students':names,'courses':courses})
    usernames=[username for cls in parsed for _,username in cls['students']]
    duplicates=[n for n,count in collections.Counter(usernames).items() if count>1]
    if duplicates:
        raise RuntimeError('Duplicate student login names across classes; resolve them before importing.')
    workbook.close()
    return parsed


def summary(parsed):
    return [{'class':c['name'],'students':len(c['students']),'courses':[course['title'] for course in c['courses']]} for c in parsed]


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook',type=Path)
    parser.add_argument('--apply',action='store_true')
    parser.add_argument('--password-stdin',action='store_true')
    args=parser.parse_args()
    parsed=parse_workbook(args.workbook)
    print(json.dumps({'workbook_summary':summary(parsed),'total_students':sum(len(c['students']) for c in parsed)},ensure_ascii=False),flush=True)
    if not args.apply:
        return
    if not args.password_stdin:
        raise RuntimeError('Provide the authorized initial password on stdin with --password-stdin.')
    password=sys.stdin.readline().rstrip('\r\n')
    if not password or len(password)>200:
        raise RuntimeError('Initial password is empty or too long.')
    spec=importlib.util.spec_from_file_location('cf_deploy',ROOT/'scripts/deploy-cloudflare.py')
    cf=importlib.util.module_from_spec(spec);spec.loader.exec_module(cf)
    deploy=ROOT/'cloudflare/wrangler.deploy.json'
    if not deploy.exists():
        raise RuntimeError('Missing existing deployment config; restore its resource IDs before importing.')
    config=json.loads(deploy.read_text())
    if config.get('account_id')!=cf.ACCOUNT:
        raise RuntimeError('Deployment config and credential account do not match.')
    database=config['d1_databases'][0]['database_id']
    endpoint='/accounts/'+cf.ACCOUNT+'/d1/database/'+database+'/query'
    def query(sql,*params):
        result=cf.api(endpoint,'POST',{'sql':sql,'params':list(params)})
        if not result or any(not item.get('success',True) for item in result):
            raise RuntimeError('Database operation failed; rerun after diagnosing the error.')
        return result[0]
    masters=query("SELECT id FROM users WHERE role='teacher' AND demo=0")['results']
    if len(masters)!=1:
        raise RuntimeError('Initialize exactly one master account first.')
    owner=masters[0]['id']
    classes=query('SELECT id,name,program FROM classes WHERE owner_id=?',owner)['results']
    if len(classes)!=5:
        raise RuntimeError('Expected the existing five class shells; preserve other classes and reconcile before import.')
    for entry in parsed:
        candidates=[c for c in classes if c['program']==entry['program']]
        if entry['number'] is not None:
            candidates=[c for c in candidates if re.search(r'NCUK\s*'+str(entry['number'])+r'\s*班',c['name'],re.I)]
        if len(candidates)!=1:
            raise RuntimeError('Cannot match an existing class safely for '+entry['sheet'])
        entry['id']=candidates[0]['id']
    if len({c['id'] for c in parsed})!=5:
        raise RuntimeError('A class was matched more than once; no changes made.')
    columns={c['name'] for c in query('PRAGMA table_info(users)')['results']}
    if 'initial_password' not in columns:
        raise RuntimeError('Apply the student account/course migration before importing.')
    query('SELECT count(*) FROM class_courses')
    existing=query("SELECT id,name,username,class_id,password IS NOT NULL AS ready FROM users WHERE owner_id=? AND role='student'",owner)['results']
    roster_names=[u for c in parsed for _,u in c['students']]
    occupied=query('SELECT id,username FROM users WHERE username IN (SELECT value FROM json_each(?))',json.dumps(roster_names,ensure_ascii=False))['results']
    occupied={u['username']:u['id'] for u in occupied}
    additions=[];activations=[];now=int(time.time()*1000);preserved=0
    for entry in parsed:
        for name,username in entry['students']:
            matches=[s for s in existing if s['class_id']==entry['id'] and s['name']==name]
            if len(matches)>1:
                raise RuntimeError('Multiple existing student records match; no changes made.')
            old=matches[0] if matches else None
            if username in occupied and (not old or occupied[username]!=old['id']):
                raise RuntimeError('A requested login name is already assigned to another account; no changes made.')
            if old and old['ready']:
                if old['username']!=username:
                    raise RuntimeError('An existing activated login differs from the roster name; preserve it until reconciled.')
                preserved+=1;continue
            salt=os.urandom(16)
            digest=hashlib.pbkdf2_hmac('sha256',password.encode(),salt,100000).hex()
            stored='pbkdf2-sha256:100000:'+salt.hex()+':'+digest
            if old:
                activations.append([old['id'],username,stored])
            else:
                additions.append([str(uuid.uuid4()),'student',name,stored,owner,entry['program'],entry['id'],0,now,username,1])
    # All input and namespace validation precedes writes. Each bulk operation is
    # atomic; reruns preserve activated passwords and stable historical IDs.
    class_rows=[{'id':c['id'],'name':c['name']} for c in parsed]
    encoded=json.dumps(class_rows,ensure_ascii=False)
    query("UPDATE classes SET name=(SELECT json_extract(value,'$.name') FROM json_each(?) WHERE json_extract(value,'$.id')=classes.id) WHERE owner_id=? AND id IN (SELECT json_extract(value,'$.id') FROM json_each(?))",encoded,owner,encoded)
    course_rows=[[c['id'],course['subject'],course['title'],course['content'],course['source']] for c in parsed for course in c['courses']]
    # This import adds/reconciles the documented courses. It does not delete any
    # separately added master courses or overwrite a teacher-edited outline.
    query("INSERT INTO class_courses(class_id,subject,title,content,source) SELECT json_extract(value,'$[0]'),json_extract(value,'$[1]'),json_extract(value,'$[2]'),json_extract(value,'$[3]'),json_extract(value,'$[4]') FROM json_each(?) WHERE true ON CONFLICT(class_id,subject) DO UPDATE SET title=excluded.title,content=excluded.content,source=excluded.source WHERE class_courses.source=excluded.source",json.dumps(course_rows,ensure_ascii=False))
    if additions:
        query("INSERT INTO users(id,role,name,password,owner_id,program,class_id,demo,created_at,username,initial_password) SELECT json_extract(value,'$[0]'),json_extract(value,'$[1]'),json_extract(value,'$[2]'),json_extract(value,'$[3]'),json_extract(value,'$[4]'),json_extract(value,'$[5]'),json_extract(value,'$[6]'),json_extract(value,'$[7]'),json_extract(value,'$[8]'),json_extract(value,'$[9]'),json_extract(value,'$[10]') FROM json_each(?)",json.dumps(additions,ensure_ascii=False))
    for student_id,username,stored in activations:
        query('UPDATE users SET username=?,password=?,invite_hash=NULL,initial_password=1 WHERE id=? AND owner_id=? AND password IS NULL',username,stored,student_id,owner)
    actual=query("SELECT c.name,count(u.id) AS students,sum(u.password IS NOT NULL) AS accounts FROM classes c LEFT JOIN users u ON u.class_id=c.id AND u.role='student' WHERE c.owner_id=? GROUP BY c.id ORDER BY c.created_at",owner)['results']
    print(json.dumps({'created_accounts':len(additions),'activated_existing':len(activations),'preserved_accounts':preserved,'class_summary':actual},ensure_ascii=False),flush=True)
    # Private audit stores source fingerprint and counts only, never passwords,
    # hashes, session cookies or individual feedback/performance content.
    directory=ROOT/'data/imports';directory.mkdir(parents=True,exist_ok=True,mode=0o700)
    audit={'source_sha256':hashlib.sha256(args.workbook.read_bytes()).hexdigest(),'imported_at':now,'summary':actual,'courses':summary(parsed)}
    descriptor=os.open(directory/'feedback-import.json',os.O_WRONLY|os.O_CREAT|os.O_TRUNC,0o600)
    with os.fdopen(descriptor,'w') as output:
        json.dump(audit,output,ensure_ascii=False,indent=2)


if __name__=='__main__':
    try:
        main()
    except RuntimeError as error:
        sys.exit(str(error))
